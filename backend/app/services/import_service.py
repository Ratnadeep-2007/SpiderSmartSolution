import csv
import io
import uuid
from typing import List, Dict, Any, Tuple, Optional
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from ..schemas.record import RecordCreate
from ..models.master import Entity, Department, EntityType
from .record_service import create_record

from openpyxl import load_workbook

async def parse_csv_preview(content: str) -> Tuple[List[str], List[List[str]]]:
    """Parses first few lines of CSV for mapping preview."""
    f = io.StringIO(content)
    reader = csv.reader(f)
    headers = next(reader, [])
    rows = []
    for _ in range(5):
        try:
            rows.append(next(reader))
        except StopIteration:
            break
    return headers, rows

async def parse_xlsx_preview(file_content: bytes) -> Tuple[List[str], List[List[str]]]:
    """Parses first few lines of XLSX for mapping preview."""
    wb = load_workbook(filename=io.BytesIO(file_content), read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(max_row=6, values_only=True):
        rows.append([str(cell) if cell is not None else "" for cell in row])
    
    if not rows:
        return [], []
    
    headers = rows[0]
    preview_rows = rows[1:]
    return headers, preview_rows

async def process_import(
    db: AsyncSession, 
    content: str, 
    field_map: Dict[str, str], 
    defaults: Dict[str, Any],
    user_id: uuid.UUID,
    is_xlsx: bool = False,
    file_bytes: Optional[bytes] = None
) -> Dict[str, Any]:
    """Processes the full CSV or XLSX import."""
    success_count = 0
    errors = []
    
    if is_xlsx and file_bytes:
        wb = load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return {"success_count": 0, "error_count": 0, "errors": []}
        
        headers = rows[0]
        data_rows = rows[1:]
        
        for i, row in enumerate(data_rows):
            try:
                # Map row tuple to dict
                row_dict = {headers[j]: row[j] for j in range(len(headers)) if j < len(row)}
                await _import_row(db, row_dict, field_map, defaults, user_id)
                success_count += 1
            except Exception as e:
                await db.rollback()
                errors.append({"row": i + 2, "error": str(e)})
    else:
        f = io.StringIO(content)
        reader = csv.DictReader(f)
        for i, row in enumerate(reader):
            try:
                await _import_row(db, row, field_map, defaults, user_id)
                success_count += 1
            except Exception as e:
                await db.rollback()
                errors.append({"row": i + 2, "error": str(e)})
            
    return {
        "success_count": success_count,
        "error_count": len(errors),
        "errors": errors
    }

async def _import_row(db, row, field_map, defaults, user_id):
    # Construct record data
    record_data = defaults.copy()
    for record_field, csv_header in field_map.items():
        if csv_header in row:
            val = row[csv_header]
            if val is None or val == "": continue
            record_data[record_field] = val
    
    # Special Handling for Relational Fields (Resolve Names to IDs if needed)
    if "entity" in record_data and "entity_id" not in record_data:
        res = await db.execute(select(Entity).where(Entity.name == record_data["entity"]))
        ent = res.scalar_one_or_none()
        if ent: record_data["entity_id"] = ent.id

    if "department" in record_data and "department_id" not in record_data:
        res = await db.execute(select(Department).where(Department.name == record_data["department"]))
        dept = res.scalar_one_or_none()
        if dept: record_data["department_id"] = dept.id

    if "entity_type" in record_data and "entity_type_id" not in record_data:
        res = await db.execute(select(EntityType).where(EntityType.name == record_data["entity_type"]))
        et = res.scalar_one_or_none()
        if et: record_data["entity_type_id"] = et.id

    # Fallback to the first available EntityType if not resolved
    if "entity_type_id" not in record_data or not record_data["entity_type_id"]:
        res = await db.execute(select(EntityType).limit(1))
        et = res.scalar_one_or_none()
        if et: record_data["entity_type_id"] = et.id

    # Validate and create
    record_in = RecordCreate(**record_data)
    await create_record(db, record_in, user_id)
