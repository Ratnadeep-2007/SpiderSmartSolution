import io
import uuid
import csv
from datetime import datetime, date
from typing import List, Any, Dict, Optional
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from ..models.record import InventoryRecord

# ReportLab imports for pure-Python, cross-platform PDF generation
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class NumberedCanvas(canvas.Canvas):
    """
    Two-pass canvas to dynamically compute and render "Page X of Y" footers,
    along with consistent headers and layout borders.
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#64748b"))
        
        # Header (rendered on pages after the first page)
        if self._pageNumber > 1:
            self.drawString(54, 750, "SpiderSmart Inventory Management System")
            self.setStrokeColor(colors.HexColor("#cbd5e1"))
            self.setLineWidth(0.5)
            self.line(54, 742, 558, 742)
            
        # Footer
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 40, page_text)
        self.drawString(54, 40, "Confidential - Internal Use Only")
        self.setStrokeColor(colors.HexColor("#cbd5e1"))
        self.setLineWidth(0.5)
        self.line(54, 52, 558, 52)
        
        self.restoreState()


def build_pdf_in_memory(title: str, headers: List[str], rows: List[List[Any]], col_widths: List[float] = None) -> io.BytesIO:
    """
    Renders table data to a beautifully formatted PDF document in memory.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=72,
        bottomMargin=72
    )
    
    styles = getSampleStyleSheet()
    
    # Custom, premium styling tokens
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=8
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=20
    )
    
    table_header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#1e293b")
    )
    
    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#334155")
    )
    
    table_cell_mono = ParagraphStyle(
        'TableCellMono',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#0f172a")
    )
    
    story = []
    
    # Title & Subtitle block
    story.append(Paragraph(title, title_style))
    story.append(Paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", subtitle_style))
    
    table_data = []
    
    # Add Header Row
    header_row = [Paragraph(h, table_header_style) for h in headers]
    table_data.append(header_row)
    
    # Add Content Rows with column styles
    for row in rows:
        formatted_row = []
        for i, val in enumerate(row):
            val_str = str(val) if val is not None else ""
            # Apply monospace styling to barcodes/identifiers for clean reading
            if i in [0, 1] and len(val_str) > 0 and any(k in headers[i].lower() for k in ["barcode", "id"]):
                style = table_cell_mono
            else:
                style = table_cell_style
            formatted_row.append(Paragraph(val_str, style))
        table_data.append(formatted_row)
        
    # Build the styled table
    t = Table(table_data, colWidths=col_widths)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#f8fafc")),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 0), (-1, 0), 6),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ('LINEBELOW', (0, 0), (-1, 0), 1, colors.HexColor("#cbd5e1")),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
        ('TOPPADDING', (0, 1), (-1, -1), 5),
    ]))
    
    story.append(t)
    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


async def generate_records_pdf(db: AsyncSession, record_ids: List[uuid.UUID]) -> io.BytesIO:
    """
    Generates a formatted PDF for a list of records.
    Returns a PDF buffer generated via ReportLab.
    """
    result = await db.execute(
        select(InventoryRecord).where(InventoryRecord.id.in_(record_ids))
    )
    records = result.scalars().all()

    headers = ["Box Barcode", "File Barcode", "Entity", "Department", "Date", "Status"]
    rows = []
    for r in records:
        r_date = r.record_date.strftime("%Y-%m-%d") if isinstance(r.record_date, (date, datetime)) else str(r.record_date)
        rows.append([
            r.box_barcode or "",
            r.file_barcode or "",
            r.entity or "",
            r.department or "",
            r_date or "",
            r.disposition_status or ""
        ])
        
    # Standard printable width is 504 points (612 page width - 54 left margin - 54 right margin)
    col_widths = [80, 80, 100, 100, 74, 70]
    return build_pdf_in_memory("SpiderSmart Inventory Records Report", headers, rows, col_widths)


async def generate_records_xlsx(db: AsyncSession, record_ids: List[uuid.UUID]) -> io.BytesIO:
    """
    Generates an Excel spreadsheet for a list of records.
    """
    if not OPENPYXL_AVAILABLE:
        buffer = io.BytesIO()
        buffer.write(b"Error: openpyxl library missing.\n")
        buffer.seek(0)
        return buffer

    result = await db.execute(
        select(InventoryRecord).where(InventoryRecord.id.in_(record_ids))
    )
    records = result.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory Records"

    # Header
    headers = ["ID", "Box Barcode", "File Barcode", "Entity", "Department", "Location", "Date", "Status", "Description"]
    ws.append(headers)

    # Style header
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
        cell.fill = openpyxl.styles.PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")

    for r in records:
        ws.append([
            str(r.id),
            r.box_barcode,
            r.file_barcode,
            r.entity,
            r.department,
            r.location,
            str(r.record_date),
            r.disposition_status,
            r.description
        ])

    # Adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = min(adjusted_width, 50)

    xlsx_buffer = io.BytesIO()
    wb.save(xlsx_buffer)
    xlsx_buffer.seek(0)
    return xlsx_buffer


async def generate_report_export(
    report_type: str, 
    data: List[Dict[str, Any]], 
    format: str,
    custom_cols: List[str] = None
) -> tuple[io.BytesIO, str, str]:
    """
    Generates a CSV or PDF for a specific report data.
    Returns (buffer, media_type, filename).
    """
    # Determine Title and Headers based on report category
    if report_type == "entity":
        title = "Records by Entity"
        headers = ["Entity", "Record Count"]
        rows = [[r.get("name") or "Unknown", r.get("value") or 0] for r in data]
        col_widths = [300, 204]
    elif report_type == "dept":
        title = "Records by Department"
        headers = ["Department", "Record Count"]
        rows = [[r.get("name") or "Unknown", r.get("value") or 0] for r in data]
        col_widths = [300, 204]
    elif report_type == "year":
        title = "Records by Year"
        headers = ["Year", "Record Count"]
        rows = [[r.get("name") or "Unknown", r.get("value") or 0] for r in data]
        col_widths = [300, 204]
    elif report_type == "location":
        title = "Records by Location"
        headers = ["Location", "Record Count"]
        rows = [[r.get("name") or "Unknown", r.get("value") or 0] for r in data]
        col_widths = [300, 204]
    elif report_type == "compliance":
        title = "Retention Compliance"
        headers = ["Status", "Record Count"]
        rows = [[r.get("name") or "Unknown", r.get("value") or 0] for r in data]
        col_widths = [300, 204]
    elif report_type == "user":
        title = "Activity by User"
        headers = ["User Email", "Record Count"]
        rows = [[r.get("name") or "Unknown", r.get("value") or 0] for r in data]
        col_widths = [300, 204]
    elif report_type == "upcoming":
        title = "Upcoming Dispositions"
        headers = ["Box Barcode", "File Barcode", "Due Date", "Entity"]
        rows = [[r.get("box_barcode") or "", r.get("file_barcode") or "", r.get("due_date") or "", r.get("entity") or ""] for r in data]
        col_widths = [120, 120, 120, 144]
    elif report_type == "holds":
        title = "Active Legal Holds"
        headers = ["Box Barcode", "File Barcode", "Entity"]
        rows = [[r.get("box_barcode") or "", r.get("file_barcode") or "", r.get("entity") or ""] for r in data]
        col_widths = [150, 150, 204]
    elif report_type == "custom":
        title = "Custom Report"
        cols = custom_cols or []
        headers = [c.replace('_', ' ').title() for c in cols]
        rows = [[r.get(c) for c in cols] for r in data]
        num_cols = len(headers)
        col_widths = [504 / num_cols] * num_cols if num_cols > 0 else None
    else:
        title = "Inventory Report"
        headers = ["Data"]
        rows = [[str(r)] for r in data]
        col_widths = [504]

    # Handle empty data gracefully
    if not rows:
        rows = [["No data available for this report"] + [""] * (len(headers) - 1)]

    # Format output properties
    filename_base = title.lower().replace(" ", "_")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    if format == "pdf":
        buffer = build_pdf_in_memory(title, headers, rows, col_widths)
        return buffer, "application/pdf", f"{filename_base}_{timestamp}.pdf"
    
    # CSV format fallback
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    
    buffer = io.BytesIO(output.getvalue().encode('utf-8'))
    buffer.seek(0)
    return buffer, "text/csv", f"{filename_base}_{timestamp}.csv"
